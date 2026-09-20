import os
import re

from docx import Document
from flask import Flask, jsonify, render_template, request
from flask_restx import Api, Resource, fields
from flask_sqlalchemy import SQLAlchemy

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv(
    "DATABASE_URL",
    "postgresql+psycopg2://cvbuilder:cvbuilder@localhost/cvbuilder",
)
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False


@app.route("/")
def index():
    return render_template("index.html")


api = Api(
    app,
    title="CV Builder API",
    version="1.0",
    description="REST API for managing CV entries with PostgreSQL.",
    doc="/docs",
)

db = SQLAlchemy(app)


class CvEntry(db.Model):
    __tablename__ = "cv_entries"

    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    opgave = db.Column(db.Text, nullable=False)
    rolle = db.Column(db.String(200), nullable=False)
    resultat = db.Column(db.Text, nullable=False)

    def to_dict(self):
        return {
            "id": self.id,
            "title": self.title,
            "opgave": self.opgave,
            "rolle": self.rolle,
            "resultat": self.resultat,
        }


with app.app_context():
    db.create_all()


cv_entry_model = api.model(
    "CvEntry",
    {
        "title": fields.String(required=True, description="Title"),
        "opgave": fields.String(required=True, description="Opgave"),
        "rolle": fields.String(required=True, description="Rolle"),
        "resultat": fields.String(required=True, description="Resultat"),
    },
)


FIELD_ALIASES = {
    "title": {"title", "titel", "projekt", "projektnavn", "navn", "overskrift"},
    "opgave": {"opgave", "task", "assignment", "arbejdsopgave", "beskrivelse", "scope", "aktivitet"},
    "rolle": {"rolle", "role", "position", "jobtitle", "ansvar", "funktion", "stilling"},
    "resultat": {"resultat", "result", "outcome", "effekt", "impact", "udbytte", "maalsaetning"},
}


def normalize_label(value):
    cleaned = (value or "").lower()
    replacements = {
        "æ": "ae",
        "ø": "oe",
        "å": "aa",
        "é": "e",
        "è": "e",
        "ü": "u",
        "ö": "o",
        "á": "a",
    }
    for source, target in replacements.items():
        cleaned = cleaned.replace(source, target)
    return re.sub(r"[^a-z0-9]+", "", cleaned)


def field_for_header(header):
    normalized = normalize_label(header)
    for canonical, aliases in FIELD_ALIASES.items():
        if normalized in {normalize_label(alias) for alias in aliases}:
            return canonical
    return None


def detect_field_indices(headers):
    index_map = {}
    for index, header in enumerate(headers):
        canonical = field_for_header(header)
        if canonical:
            index_map[canonical] = index
    return index_map


def table_to_rows(table):
    if not table.rows:
        return []

    headers = [cell.text.strip() for cell in table.rows[0].cells]
    if not headers:
        return []

    index_map = detect_field_indices(headers)
    if not index_map:
        return []

    rows = []
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        if not any(cells):
            continue

        if len(cells) <= max(index_map.values()):
            continue

        payload = {}
        for field_name, index in index_map.items():
            payload[field_name] = cells[index]

        if all(payload.get(field_name, "").strip() for field_name in ["title", "opgave", "rolle", "resultat"]):
            rows.append(payload)

    return rows


def table_to_rows_fallback(table):
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        values = [cell for cell in cells if cell]
        if len(values) < 4:
            continue
        candidate = {
            "title": values[0],
            "opgave": values[1] if len(values) > 1 else "",
            "rolle": values[2] if len(values) > 2 else "",
            "resultat": values[3] if len(values) > 3 else "",
        }
        if all(candidate.get(field, "").strip() for field in ["title", "opgave", "rolle", "resultat"]):
            rows.append(candidate)
    return rows


def extract_docx_rows(file_obj):
    document = Document(file_obj)
    extracted = []

    for table in document.tables:
        rows = table_to_rows(table)
        if rows:
            extracted.extend(rows)
        elif len(table.rows) >= 2:
            fallback_rows = table_to_rows_fallback(table)
            if fallback_rows:
                extracted.extend(fallback_rows)

    if extracted:
        return extracted

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if re.search(r"(title|titel|opgave|rolle|resultat|result)", text, flags=re.I):
            if ":" in text:
                pair = text.split(":", 1)
                if len(pair) == 2 and pair[0].strip():
                    extracted.append({
                        "title": pair[0].strip() if re.search(r"(title|titel)", pair[0], flags=re.I) else "",
                        "opgave": pair[1].strip() if re.search(r"(opgave|task)", pair[0], flags=re.I) else "",
                        "rolle": pair[1].strip() if re.search(r"(rolle|role)", pair[0], flags=re.I) else "",
                        "resultat": pair[1].strip() if re.search(r"(resultat|result|outcome)", pair[0], flags=re.I) else "",
                    })

    return [entry for entry in extracted if all(entry.get(field, "").strip() for field in ["title", "opgave", "rolle", "resultat"])]


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


@api.route("/api/items")
class CvEntryList(Resource):
    @api.doc("list_cv_entries")
    def get(self):
        items = CvEntry.query.order_by(CvEntry.id).all()
        return [item.to_dict() for item in items]

    @api.doc("create_cv_entry")
    @api.expect(cv_entry_model, validate=True)
    def post(self):
        payload = api.payload or {}
        required = ["title", "opgave", "rolle", "resultat"]

        for key in required:
            if not payload.get(key):
                return {"message": f"Missing required field: {key}"}, 400

        entry = CvEntry(
            title=payload["title"],
            opgave=payload["opgave"],
            rolle=payload["rolle"],
            resultat=payload["resultat"],
        )
        db.session.add(entry)
        db.session.commit()
        return entry.to_dict(), 201


@api.route("/api/items/<int:item_id>")
class CvEntryDetail(Resource):
    @api.doc("get_cv_entry")
    def get(self, item_id):
        item = db.session.get(CvEntry, item_id)
        if not item:
            return {"message": "Item not found"}, 404
        return item.to_dict()

    @api.doc("update_cv_entry")
    @api.expect(cv_entry_model, validate=False)
    def put(self, item_id):
        item = db.session.get(CvEntry, item_id)
        if not item:
            return {"message": "Item not found"}, 404

        data = api.payload or {}
        if "title" in data and data["title"]:
            item.title = data["title"]
        if "opgave" in data and data["opgave"]:
            item.opgave = data["opgave"]
        if "rolle" in data and data["rolle"]:
            item.rolle = data["rolle"]
        if "resultat" in data and data["resultat"]:
            item.resultat = data["resultat"]

        db.session.commit()
        return item.to_dict()

    @api.doc("delete_cv_entry")
    def delete(self, item_id):
        item = db.session.get(CvEntry, item_id)
        if not item:
            return {"message": "Item not found"}, 404

        db.session.delete(item)
        db.session.commit()
        return {"message": "Item deleted"}, 200


@app.route("/api/docx/import", methods=["POST"])
def import_docx():
    uploaded = request.files.get("file")
    if not uploaded:
        return jsonify({"message": "No file uploaded"}), 400

    filename = (uploaded.filename or "").lower()
    if not filename.endswith(".docx"):
        return jsonify({"message": "Only .docx files are supported"}), 400

    try:
        rows = extract_docx_rows(uploaded)
    except Exception as exc:  # pragma: no cover - defensive error for uploaded file
        return jsonify({"message": f"Could not read DOCX file: {exc}"}), 400

    if not rows:
        return jsonify({"message": "No matching table rows with Title/Opgave/Rolle/Resultat were found."}), 400

    return jsonify({"rows": rows, "count": len(rows)})


@app.route("/api/docx/save", methods=["POST"])
def save_docx_entries():
    data = request.get_json(silent=True) or {}
    rows = data.get("rows") or []
    if not isinstance(rows, list) or not rows:
        return jsonify({"message": "No rows supplied"}), 400

    saved = []
    for row in rows:
        required = ["title", "opgave", "rolle", "resultat"]
        if not all(str(row.get(field, "")).strip() for field in required):
            continue

        entry = CvEntry(
            title=row["title"].strip(),
            opgave=row["opgave"].strip(),
            rolle=row["rolle"].strip(),
            resultat=row["resultat"].strip(),
        )
        db.session.add(entry)
        saved.append(entry)

    if not saved:
        return jsonify({"message": "No valid rows were saved"}), 400

    db.session.commit()
    return jsonify({"saved": len(saved), "items": [item.to_dict() for item in saved]})


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
