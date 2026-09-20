import os
import tempfile
import unittest
from io import BytesIO

from docx import Document

from app import create_app
from cv_builder.extensions import db
from cv_builder.models import CvEntry


class TestConfig:
    TESTING = True
    SQLALCHEMY_TRACK_MODIFICATIONS = False


class CvBuilderTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.database_file = tempfile.NamedTemporaryFile(suffix=".sqlite", delete=False)
        cls.database_file.close()
        TestConfig.SQLALCHEMY_DATABASE_URI = f"sqlite:///{cls.database_file.name}"
        cls.app = create_app(TestConfig)
        cls.client = cls.app.test_client()

    @classmethod
    def tearDownClass(cls):
        with cls.app.app_context():
            db.session.remove()
            db.engine.dispose()
        os.unlink(cls.database_file.name)

    def setUp(self):
        with self.app.app_context():
            db.session.query(CvEntry).delete()
            db.session.commit()

    @staticmethod
    def entry_payload(title="Project A"):
        return {
            "title": title,
            "rolle": "Developer",
            "opgave": "Build a feature",
            "resultat": "Released successfully",
        }

    def test_pages_and_health_are_available(self):
        for path in ["/", "/file-upload", "/manual-entry", "/new-entry", "/deleted-entries"]:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

        self.assertEqual(self.client.get("/health").get_json(), {"status": "ok"})

    def test_entry_crud_and_lock_protection(self):
        created = self.client.post("/api/items/all", json=self.entry_payload())
        self.assertEqual(created.status_code, 201)
        item_id = created.get_json()["id"]

        locked = self.client.put(f"/api/items/{item_id}", json={"locked": True})
        self.assertEqual(locked.status_code, 200)
        self.assertTrue(locked.get_json()["locked"])

        blocked_delete = self.client.delete(f"/api/items/{item_id}")
        self.assertEqual(blocked_delete.status_code, 409)
        self.assertEqual(self.client.get(f"/api/items/{item_id}").status_code, 200)

        unlocked = self.client.put(f"/api/items/{item_id}", json={"locked": False})
        self.assertEqual(unlocked.status_code, 200)
        deleted = self.client.delete(f"/api/items/{item_id}")
        self.assertEqual(deleted.status_code, 200)

        active_items = self.client.get("/api/items/all").get_json()["items"]
        deleted_items = self.client.get("/api/items/deleted").get_json()["items"]
        self.assertFalse(any(item["id"] == item_id for item in active_items))
        self.assertTrue(any(item["id"] == item_id for item in deleted_items))

        recovered = self.client.post(f"/api/items/{item_id}/recover")
        self.assertEqual(recovered.status_code, 200)
        active_items = self.client.get("/api/items/all").get_json()["items"]
        self.assertTrue(any(item["id"] == item_id for item in active_items))

    def test_docx_import_preserves_requested_column_order(self):
        document = Document()
        table = document.add_table(rows=2, cols=4)
        headers = ["Title", "Rolle", "Opgave", "Resultat"]
        values = ["Project A", "Developer", "Build a feature", "Released successfully"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        buffer = BytesIO()
        document.save(buffer)
        response = self.client.post(
            "/api/docx/import",
            data={"file": (BytesIO(buffer.getvalue()), "cv.docx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["rows"],
            [{
                "title": "Project A",
                "rolle": "Developer",
                "opgave": "Build a feature",
                "resultat": "Released successfully",
            }],
        )


if __name__ == "__main__":
    unittest.main()
