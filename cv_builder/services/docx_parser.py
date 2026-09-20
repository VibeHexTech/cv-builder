import re

from docx import Document


FIELD_ALIASES = {
    "title": {"title", "titel", "projekt", "projektnavn", "navn", "overskrift"},
    "opgave": {"opgave", "task", "assignment", "arbejdsopgave", "beskrivelse", "scope", "aktivitet"},
    "rolle": {"rolle", "role", "position", "jobtitle", "ansvar", "funktion", "stilling"},
    "resultat": {"resultat", "result", "outcome", "effekt", "impact", "udbytte", "maalsaetning"},
}
REQUIRED_FIELDS = ("title", "opgave", "rolle", "resultat")


def normalize_label(value):
    cleaned = (value or "").lower()
    replacements = {
        "æ": "ae",
        "ø": "oe",
        "å": "aa",
        "é": "e",
        "è": "e",
        "ü": "u",
        "ö": "o",
        "á": "a",
    }
    for source, target in replacements.items():
        cleaned = cleaned.replace(source, target)
    return re.sub(r"[^a-z0-9]+", "", cleaned)


def field_for_header(header):
    normalized = normalize_label(header)
    for canonical, aliases in FIELD_ALIASES.items():
        if normalized in {normalize_label(alias) for alias in aliases}:
            return canonical
    return None


def detect_field_indices(headers):
    index_map = {}
    for index, header in enumerate(headers):
        canonical = field_for_header(header)
        if canonical:
            index_map[canonical] = index
    return index_map


def table_to_rows(table):
    if not table.rows:
        return []

    headers = [cell.text.strip() for cell in table.rows[0].cells]
    if not headers:
        return []

    index_map = detect_field_indices(headers)
    if not index_map:
        return []

    rows = []
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        if not any(cells) or len(cells) <= max(index_map.values()):
            continue

        payload = {field_name: cells[index] for field_name, index in index_map.items()}
        if all(payload.get(field, "").strip() for field in REQUIRED_FIELDS):
            rows.append(payload)

    return rows


def table_to_rows_fallback(table):
    rows = []
    for row in table.rows:
        values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
        if len(values) < 4:
            continue

        candidate = {
            "title": values[0],
            "rolle": values[1],
            "opgave": values[2],
            "resultat": values[3],
        }
        if all(candidate[field].strip() for field in REQUIRED_FIELDS):
            rows.append(candidate)

    return rows


def extract_docx_rows(file_obj):
    document = Document(file_obj)
    extracted = []

    for table in document.tables:
        rows = table_to_rows(table)
        if rows:
            extracted.extend(rows)
        elif len(table.rows) >= 2:
            extracted.extend(table_to_rows_fallback(table))

    if extracted:
        return extracted

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text or not re.search(r"(title|titel|opgave|rolle|resultat|result)", text, flags=re.I):
            continue
        if ":" not in text:
            continue

        label, value = text.split(":", 1)
        label = label.strip()
        value = value.strip()
        extracted.append({
            "title": value if re.search(r"(title|titel)", label, flags=re.I) else "",
            "opgave": value if re.search(r"(opgave|task)", label, flags=re.I) else "",
            "rolle": value if re.search(r"(rolle|role)", label, flags=re.I) else "",
            "resultat": value if re.search(r"(resultat|result|outcome)", label, flags=re.I) else "",
        })

    return [entry for entry in extracted if all(entry[field].strip() for field in REQUIRED_FIELDS)]
